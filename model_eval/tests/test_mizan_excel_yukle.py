"""mizan_excel_yukle.py::_mizan_satirlarini_oku - başlık satırının otomatik
bulunması (2026-09-30, gerçek bug'ın düzeltmesi). Önceden başlık satırı
`min_row=7` ile SABİTTİ - bu sadece Archive2/mizan.xlsx'in kendi formatına
(üstünde 5 boş satır) özgüydü. ONBOARDING-VERI-FORMATI.md dış ekibe
başlığın herhangi bir satırda olabileceğini ima ediyor ("veri, başlık
satırından sonraki satırdan itibaren okunur") - standart formatta (1. satır
başlık) bir mizan gönderilirse eski kod SESSİZCE 0 satır okuyordu, hiç
hata vermiyordu."""

import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from mizan_excel_yukle import _mizan_satirlarini_oku


def _mizan_yaz(tmp_path, baslik_satiri, veriler, sutun_sirasi=("HESAP KODU", "HESAP ADI")):
    wb = openpyxl.Workbook()
    ws = wb.active
    for _ in range(baslik_satiri - 1):
        ws.append([])
    ws.append(list(sutun_sirasi))
    for kod, ad in veriler:
        satir = [None, None]
        satir[sutun_sirasi.index("HESAP KODU")] = kod
        satir[sutun_sirasi.index("HESAP ADI")] = ad
        ws.append(satir)
    yol = tmp_path / "mizan.xlsx"
    wb.save(yol)
    return yol


class TestBaslikSatiriOtomatikBulma:
    def test_standart_format_1_satir_baslik(self, tmp_path):
        """Dış ekibin göndermesi beklenen standart format: başlık 1. satır."""
        yol = _mizan_yaz(tmp_path, 1, [("120.01.00001", "Test Müşteri")])
        sonuc = _mizan_satirlarini_oku(yol)
        assert sonuc == [("120.01.00001", "Test Müşteri")]

    def test_gercek_format_6_satir_baslik(self, tmp_path):
        """Archive2/mizan.xlsx'in kendi formatı - üstünde 5 boş satır."""
        yol = _mizan_yaz(tmp_path, 6, [("320.01.00001", "Test Tedarikçi")])
        sonuc = _mizan_satirlarini_oku(yol)
        assert sonuc == [("320.01.00001", "Test Tedarikçi")]

    def test_sutun_sirasi_ters_olsa_da_isme_gore_bulunur(self, tmp_path):
        """ONBOARDING-VERI-FORMATI.md sütun SIRASINI belirtmiyor - isme göre
        bulunmalı, pozisyona göre değil."""
        yol = _mizan_yaz(tmp_path, 1, [("340.01.00001", "Ters Sıra Testi")],
                          sutun_sirasi=("HESAP ADI", "HESAP KODU"))
        sonuc = _mizan_satirlarini_oku(yol)
        assert sonuc == [("340.01.00001", "Ters Sıra Testi")]

    def test_baslik_hic_bulunamazsa_acik_hata(self, tmp_path):
        """Sessizce 0 satır okumak yerine açık SystemExit fırlatmalı."""
        wb = openpyxl.Workbook()
        wb.active.append(["Yanlış Başlık", "Başka Sütun"])
        wb.active.append(["120.01.00001", "Bu satır hiç okunmamalı"])
        yol = tmp_path / "mizan.xlsx"
        wb.save(yol)
        with pytest.raises(SystemExit, match="HESAP KODU"):
            _mizan_satirlarini_oku(yol)

    def test_20_satirdan_sonraki_baslik_bulunamaz(self, tmp_path):
        """Makul bir arama sınırı var - normal bir mizanda başlık bundan
        aşağıda olmamalı, olursa sessizce yanlış okumak yerine açık hata."""
        yol = _mizan_yaz(tmp_path, 25, [("120.01.00001", "Çok Aşağıda")])
        with pytest.raises(SystemExit, match="HESAP KODU"):
            _mizan_satirlarini_oku(yol)


class TestSayisalHesapKoduAtlanir:
    def test_sayisal_kod_atlanir_metin_kod_kalir(self, tmp_path):
        """HESAP KODU hücresi Excel'de sayı biçimindeyse (örn. 320 metin
        değil sayı olarak girilmişse) bu satır sessizce kaybolmamalı -
        atlanır ama diğer, metin biçimindeki satırlar etkilenmez."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["HESAP KODU", "HESAP ADI"])
        ws.append(["120.01.00001", "Metin Kod"])
        ws.append([320, "Sayısal Kod"])
        yol = tmp_path / "mizan.xlsx"
        wb.save(yol)

        sonuc = _mizan_satirlarini_oku(yol)
        assert sonuc == [("120.01.00001", "Metin Kod")]

    def test_sayisal_kod_atlanirken_uyari_loglanir(self, tmp_path, caplog):
        """Sessiz veri kaybı olmamalı - hangi satırın neden atlandığı loglanır."""
        import logging

        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["HESAP KODU", "HESAP ADI"])
        ws.append([320, "Sayısal Kod"])
        yol = tmp_path / "mizan.xlsx"
        wb.save(yol)

        with caplog.at_level(logging.WARNING):
            _mizan_satirlarini_oku(yol)
        assert any("hucresi metin degil" in r.message for r in caplog.records)

    def test_tamamen_bos_satir_sessizce_atlanir_uyarisiz(self, tmp_path, caplog):
        """Gerçekten boş bir satır (kod hiç girilmemiş) normal bir durum -
        bu, sayısal-tip hatasından ayrı ele alınmalı, uyarı loglamamalı."""
        import logging

        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["HESAP KODU", "HESAP ADI"])
        ws.append([None, None])
        ws.append(["120.01.00001", "Geçerli Kod"])
        yol = tmp_path / "mizan.xlsx"
        wb.save(yol)

        with caplog.at_level(logging.WARNING):
            sonuc = _mizan_satirlarini_oku(yol)
        assert sonuc == [("120.01.00001", "Geçerli Kod")]
        assert not any("hucresi metin degil" in r.message for r in caplog.records)


# ---------------------------------------------------------------------------
# mizan_oku — her derinlikteki yaprak hesap + sebepli atlama raporu (2026-10-19)
# Gerçek bulgu: bir müşterinin mizanında `600.01`, `191.01` gibi 2 parçalı hesaplar vardı; eski süzgeç yalnızca tam 3 parçalı
# kodları aldığı için bunlar SESSİZCE atılıyor, "mizan 17 satır" tek başına eksik mizanı gizliyordu.
# ---------------------------------------------------------------------------

from mizan_excel_yukle import mizan_hesap_satirlari, mizan_oku  # noqa: E402


def _oku(tmp_path, veriler):
    return mizan_oku(_mizan_yaz(tmp_path, 1, veriler))


class TestMizanOku:
    def test_iki_parcali_yaprak_hesaplar_alinir(self, tmp_path):
        """ANA BUG: altında alt hesap olmayan 2 parçalı hesaplar (600.01, 191.01) yapraktır ve alınmalıdır."""
        satirlar, rapor = _oku(tmp_path, [("600.01", "Tavan Kliması"), ("600.02", "Araç Buzdolabı"), ("191.01", "%20 İnd. KDV")])
        assert satirlar == [("600.01", "600", "Tavan Kliması"), ("600.02", "600", "Araç Buzdolabı"), ("191.01", "191", "%20 İnd. KDV")]
        assert rapor["derinlik"] == {"2": 3} and rapor["kabul_edilen"] == 3

    def test_ust_grup_atlanir_yaprak_alinir(self, tmp_path):
        """`120.01` altında `120.01.001` varsa `120.01` bir gruptur (alt kırılım seçeneği değil); yaprak alınır."""
        satirlar, rapor = _oku(tmp_path, [("120", "Alıcılar"), ("120.01", "Yurtiçi"), ("120.01.001", "Ali Veli"), ("120.01.002", "Ayşe")])
        assert [k for k, _, _ in satirlar] == ["120.01.001", "120.01.002"]
        assert rapor["atlanan"]["ana_kod"] == 1 and rapor["atlanan"]["ust_grup"] == 1
        assert rapor["atlanan_ornek"]["ust_grup"] == ["120.01"]

    def test_dort_parcali_ve_karisik_dolgu_genislikleri(self, tmp_path):
        """Gerçek müşteri biçimi: `150.01.001` ve `150.01.03` (farklı dolgu) aynı mizanda; 4 parçalı yaprak da alınır."""
        satirlar, rapor = _oku(tmp_path, [("150.01.001", "A"), ("150.01.03", "B"), ("120.01.001.0001", "C"), ("120.01.001", "grup")])
        assert [k for k, _, _ in satirlar] == ["150.01.001", "150.01.03", "120.01.001.0001"]
        assert rapor["derinlik"] == {"3": 2, "4": 1}
        assert rapor["atlanan"]["ust_grup"] == 1  # 120.01.001, altında 4 parçalı hesap var

    def test_ayni_mizanda_karisik_derinlik(self, tmp_path):
        satirlar, _ = _oku(tmp_path, [("191.01", "%20 İnd. KDV"), ("320.01.00272", "Satıcı"), ("600.01", "Satış")])
        assert {k for k, _, _ in satirlar} == {"191.01", "320.01.00272", "600.01"}

    def test_ana_kod_satirlari_alinmaz_ama_sayilir(self, tmp_path):
        satirlar, rapor = _oku(tmp_path, [("320", "Satıcılar"), ("600", "Satışlar"), ("600.01", "Yurtiçi")])
        assert [k for k, _, _ in satirlar] == ["600.01"] and rapor["atlanan"]["ana_kod"] == 2

    @pytest.mark.parametrize("kotu", ["12.01", "1200.01", "ABC.01", "60.01.001"])
    def test_gecersiz_ilk_parca_atlanir_ve_raporlanir(self, tmp_path, kotu):
        satirlar, rapor = _oku(tmp_path, [(kotu, "Kötü"), ("600.01", "İyi")])
        assert [k for k, _, _ in satirlar] == ["600.01"]
        assert rapor["atlanan"]["gecersiz_kod"] == 1 and any("geçersiz" in u for u in rapor["uyari"])

    def test_ayirici_normalizasyonu_tire_ve_bosluk(self, tmp_path):
        satirlar, _ = _oku(tmp_path, [("600-01", "A"), ("191 02", "B")])
        assert [k for k, _, _ in satirlar] == ["600.01", "191.02"]

    def test_tekrar_ve_sayi_bicimli_raporlanir(self, tmp_path):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["HESAP KODU", "HESAP ADI"])
        ws.append(["600.01", "İlk"])
        ws.append(["600.01", "Tekrar"])
        ws.append([320, "Sayı"])
        yol = tmp_path / "m.xlsx"
        wb.save(yol)
        satirlar, rapor = mizan_oku(yol)
        assert satirlar == [("600.01", "600", "İlk")]
        assert rapor["atlanan"]["tekrar"] == 1 and rapor["atlanan"]["sayi_bicimli"] == 1
        assert rapor["okunan_satir"] == 3 and len(rapor["uyari"]) == 2

    def test_ek_sayfalar_okunmaz_ama_raporlanir(self, tmp_path):
        """Yalnızca ilk sayfa okunur (davranış aynı) — ama diğer sayfalar artık SESSİZ değil, uyarıda adıyla görünür."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Mizan"
        ws.append(["HESAP KODU", "HESAP ADI"])
        ws.append(["600.01", "A"])
        ws2 = wb.create_sheet("Cariler")
        ws2.append(["HESAP KODU", "HESAP ADI"])
        for i in range(1, 40):
            ws2.append([f"120.01.{i:03d}", f"Müşteri {i}"])
        yol = tmp_path / "m.xlsx"
        wb.save(yol)
        satirlar, rapor = mizan_oku(yol)
        assert [k for k, _, _ in satirlar] == ["600.01"]
        assert rapor["okunmayan_sayfalar"] == [{"ad": "Cariler", "satir": 40}]
        assert any("Cariler" in u and "İLK sayfa" in u for u in rapor["uyari"])

    def test_rapor_toplamlari_tutarli(self, tmp_path):
        """okunan_satir = alınan + atlanan toplamı — hiçbir satır sayılmadan kaybolmaz."""
        _, rapor = _oku(tmp_path, [("320", "a"), ("120.01", "g"), ("120.01.001", "y"), ("600.01", "z"), ("XX.01", "k"), ("600.01", "t")])
        assert rapor["okunan_satir"] == 6
        assert rapor["kabul_edilen"] + sum(rapor["atlanan"].values()) == rapor["okunan_satir"]

    def test_hic_hesap_yoksa_uyari_verir(self, tmp_path):
        satirlar, rapor = _oku(tmp_path, [("320", "Satıcılar"), ("600", "Satışlar")])
        assert satirlar == [] and any("alınabilecek alt hesap bulunamadı" in u for u in rapor["uyari"])

    def test_hesap_satirlari_geriye_uyumlu_imza(self, tmp_path):
        yol = _mizan_yaz(tmp_path, 1, [("600.01", "A")])
        assert mizan_hesap_satirlari(yol) == [("600.01", "600", "A")]
