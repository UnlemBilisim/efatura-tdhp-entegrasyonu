"""Onboard edilmemiş şirketin public şemaya (başka şirketin verisine)
düşmemesi — core.db.tenant_kayitli_mi / get_conn guard'ı."""

import os
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core import db
from core.constants import DEFAULT_OWN_VKN
from conftest import TEST_DATABASE_URL, requires_postgres

KAYITLI_VKN = "1111111111"
KAYITSIZ_VKN = "2222222222"


class TestVeritabanisizKurallar:
    def test_varsayilan_sirket_her_zaman_kayitli(self):
        assert db.tenant_kayitli_mi(DEFAULT_OWN_VKN) is True

    @pytest.mark.parametrize("vkn", [None, "", "12345abcde"])
    def test_gecersiz_vkn_kayitli_degil(self, vkn):
        assert db.tenant_kayitli_mi(vkn) is False


@pytest.fixture
def tenant_db():
    from psycopg2 import sql

    old_url = os.environ.get("DATABASE_URL")
    os.environ["DATABASE_URL"] = TEST_DATABASE_URL
    db.reset_pool_for_tests()
    pool = db.get_pool()
    conn = pool.getconn()
    try:
        with conn.cursor() as cur:
            cur.execute(sql.SQL("CREATE SCHEMA IF NOT EXISTS {}").format(
                sql.Identifier(f"tenant_{KAYITLI_VKN}")
            ))
            cur.execute(sql.SQL("DROP SCHEMA IF EXISTS {} CASCADE").format(
                sql.Identifier(f"tenant_{KAYITSIZ_VKN}")
            ))
            cur.execute("TRUNCATE TABLE public.mizan_alt_kirilim")
            cur.execute(
                "INSERT INTO public.mizan_alt_kirilim (hesap_kodu, ana_kod, hesap_adi) "
                "VALUES ('120.01.00008', '120', 'BASKA SIRKETIN MUSTERISI')"
            )
        conn.commit()
    finally:
        pool.putconn(conn)

    yield

    conn = pool.getconn()
    try:
        with conn.cursor() as cur:
            cur.execute("TRUNCATE TABLE public.mizan_alt_kirilim")
        conn.commit()
    finally:
        pool.putconn(conn)
    db.reset_pool_for_tests()
    if old_url is None:
        os.environ.pop("DATABASE_URL", None)
    else:
        os.environ["DATABASE_URL"] = old_url


@requires_postgres
class TestTenantIzolasyonuDB:
    def test_semasi_olan_sirket_kayitli(self, tenant_db):
        assert db.tenant_kayitli_mi(KAYITLI_VKN) is True

    def test_semasi_olmayan_sirket_kayitli_degil(self, tenant_db):
        assert db.tenant_kayitli_mi(KAYITSIZ_VKN) is False

    def test_kayitsiz_sirket_public_semaya_dusmez(self, tenant_db):
        with pytest.raises(db.TenantKayitliDegilHatasi):
            with db.get_conn(tenant_vkn=KAYITSIZ_VKN):
                pass

    def test_kayitsiz_sirket_baska_sirketin_mizanini_goremez(self, tenant_db):
        import core.mizan as mizan

        mizan.reset_mizan_cache_for_tests()
        assert mizan.get_alt_kirilimlar(KAYITSIZ_VKN) == {}

    def test_mizan_yukleme_kayitsiz_sirkete_yazmaz(self, tenant_db, tmp_path):
        import openpyxl

        import mizan_excel_yukle

        excel = tmp_path / "mizan.xlsx"
        wb = openpyxl.Workbook()
        # Başlık (HESAP KODU/HESAP ADI) olmadan _baslik_satirini_bul kendi
        # format hatasını fırlatır (2026-09-30 düzeltmesi) - bu test o
        # senaryoyu değil, kayıtsız VKN korumasını izole test ediyor, bu
        # yüzden gerçek formatta başlık + veri satırı gerekiyor.
        wb.active.cell(row=6, column=1, value="HESAP KODU")
        wb.active.cell(row=6, column=2, value="HESAP ADI")
        wb.active.cell(row=7, column=1, value="120.01.00001")
        wb.active.cell(row=7, column=2, value="Yeni Musteri")
        wb.save(excel)

        with pytest.raises(SystemExit, match="onboard"):
            mizan_excel_yukle.yukle(KAYITSIZ_VKN, str(excel), database_url=TEST_DATABASE_URL)

        with db.get_conn() as conn:
            with conn.cursor() as cur:
                cur.execute("SELECT count(*) FROM public.mizan_alt_kirilim")
                assert cur.fetchone()[0] == 1
